"""Read-only, synthetic WPS spreadsheet -> meeting brief diagnostic.

This is a deterministic script baseline and guard fixture. It is not a Motif
operator, a real research result, or evidence of saved model requests.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
DEFAULT_BOOK = ROOT / "outputs" / "wps-meeting-mock-v1" / "experiment_log.xlsx"
DEFAULT_OUTPUT = ROOT / "outputs" / "wps-meeting-mock-v1" / "reports"
FIELDS = ("run_id", "variant", "seed", "test_cases", "correct", "duration_s", "source_run")
CLAIM_FIELDS = ("run_id", "variant", "seed", "test_cases", "correct", "source_run")


def read_rows(book: Path) -> list[dict]:
    workbook = load_workbook(book, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["实验记录"]:
            raise ValueError("expected one 实验记录 worksheet")
        values = workbook["实验记录"].iter_rows(values_only=True)
        header = next(values, None)
        if tuple(header or ()) != FIELDS:
            raise ValueError("worksheet schema changed; manual mapping required")
        rows = []
        for cells in values:
            if all(value is None for value in cells):
                continue
            if len(cells) != len(FIELDS):
                raise ValueError("record width changed")
            rows.append(dict(zip(FIELDS, cells, strict=True)))
        return rows
    finally:
        workbook.close()


def validate(rows: list[dict]) -> None:
    if not rows:
        raise ValueError("no measurements")
    ids = [row["run_id"] for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate run_id")
    seeds = defaultdict(set)
    for row in rows:
        if row["variant"] not in {"baseline", "candidate"}:
            raise ValueError("unexpected variant")
        if type(row["seed"]) is not int or row["seed"] <= 0:
            raise ValueError("invalid seed")
        if type(row["test_cases"]) is not int or row["test_cases"] <= 0:
            raise ValueError("invalid denominator")
        if type(row["correct"]) is not int or not 0 <= row["correct"] <= row["test_cases"]:
            raise ValueError("invalid correct count")
        if type(row["duration_s"]) not in (int, float) or row["duration_s"] < 0:
            raise ValueError("invalid duration")
        if not isinstance(row["source_run"], str) or not row["source_run"].startswith("synthetic://"):
            raise ValueError("source provenance changed")
        if row["seed"] in seeds[row["variant"]]:
            raise ValueError("duplicate seed within variant")
        seeds[row["variant"]].add(row["seed"])
    if set(seeds) != {"baseline", "candidate"} or seeds["baseline"] != seeds["candidate"]:
        raise ValueError("variants must have paired seeds")


def apply_event(rows: list[dict], event: list[dict]) -> list[dict]:
    updated = [dict(row) for row in rows]
    by_id = {row["run_id"]: row for row in updated}
    for change in event:
        field = change["field"]
        if field not in FIELDS or field == "run_id":
            raise ValueError("event field is not a permitted measurement")
        row = by_id.get(change["run_id"])
        if row is None or row[field] != change["old"]:
            raise ValueError("event precondition failed")
        row[field] = change["new"]
    validate(updated)
    return updated


def claim_dependency(rows: list[dict]) -> list[list]:
    return [[row[field] for field in CLAIM_FIELDS] for row in sorted(rows, key=lambda item: item["run_id"])]


def metrics(rows: list[dict]) -> dict:
    totals = {variant: {"correct": 0, "test_cases": 0, "duration_s": 0.0}
              for variant in ("baseline", "candidate")}
    for row in rows:
        group = totals[row["variant"]]
        group["correct"] += row["correct"]
        group["test_cases"] += row["test_cases"]
        group["duration_s"] += row["duration_s"]
    for group in totals.values():
        group["accuracy"] = group["correct"] / group["test_cases"]
        group["duration_s"] = round(group["duration_s"], 1)
    totals["delta_percentage_points"] = round(
        100 * (totals["candidate"]["accuracy"] - totals["baseline"]["accuracy"]), 2
    )
    return totals


def make_qmd(contract: dict, scenario: str, result: dict, reusable: bool,
             source_sha256: str, snapshot_sha256: str) -> str:
    status = "沿用模拟已审阅表述；仍待真实研究者复核" if reusable else "旧表述已失效；需要新的语义判断和人工复核"
    claim = contract["approved_claim"] if reusable else "本次没有生成新的科学结论，也没有沿用旧结论。"
    base, cand = result["baseline"], result["candidate"]
    return f"""---
title: "合成实验组会简报 · {scenario}"
format:
  html:
    toc: false
---

> **合成 mock，仅用于工具与守卫诊断。** 所有运行记录均为虚构；此页不代表真实实验或研究结论。

## 本轮状态

**{status}**

{claim}

## 可复算数字

| 组别 | 正确数 / 测试数 | 汇总正确率 | 总耗时（秒） |
|:---|---:|---:|---:|
| 基线 | {base['correct']} / {base['test_cases']} | {base['accuracy']:.1%} | {base['duration_s']:.1f} |
| 候选 | {cand['correct']} / {cand['test_cases']} | {cand['accuracy']:.1%} | {cand['duration_s']:.1f} |

正确率差值：**{result['delta_percentage_points']:+.2f} 个百分点**。定义：{contract['metric']}。

## 来源与边界

- 原始表格 SHA-256：`{source_sha256}`；施加事件后的记录快照 SHA-256：`{snapshot_sha256}`。
- 事件：`{scenario}`。增量在内存中模拟，没有改写原始表格。
- `duration_s` 只进入耗时列，不支持正确率结论。
- 本页没有调用模型；“需要语义判断”表示暂停在交接点，不能算一次已节省的模型请求。
"""


def run(book: Path, output_dir: Path, scenario: str, render: bool = True) -> dict:
    contract = json.loads((HERE / "mock_contract.json").read_text(encoding="utf-8"))
    if scenario not in contract["scenarios"]:
        raise ValueError(f"unknown scenario: {scenario}")
    original = read_rows(book)
    validate(original)
    rows = apply_event(original, contract["scenarios"][scenario])
    approved = sorted(contract["approved_dependency"], key=lambda item: item[0])
    reusable = claim_dependency(rows) == approved
    result = metrics(rows)
    # The frozen mock statement must match its original numeric dependencies.
    if scenario == "base" and (not reusable or result["delta_percentage_points"] != 8.0):
        raise ValueError("mock approval no longer matches base workbook")
    source_hash = hashlib.sha256(book.read_bytes()).hexdigest()
    snapshot_hash = hashlib.sha256(json.dumps(
        sorted(rows, key=lambda item: item["run_id"]), ensure_ascii=False, sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    output_dir.mkdir(parents=True, exist_ok=True)
    qmd = output_dir / f"{scenario}.qmd"
    qmd.write_text(make_qmd(contract, scenario, result, reusable, source_hash, snapshot_hash), encoding="utf-8")
    if render:
        subprocess.run(["quarto", "render", str(qmd), "--to", "html", "--no-execute"], check=True)
    log = {
        "fixture": contract["fixture"],
        "scenario": scenario,
        "source_sha256": source_hash,
        "effective_snapshot_sha256": snapshot_hash,
        "event": contract["scenarios"][scenario],
        "metric": result,
        "claim_dependency_unchanged": reusable,
        "decision": "reuse_mock_reviewed_claim" if reusable else "semantic_handoff_required",
        "model_requests_executed": 0,
        "estimated_model_requests_saved": None,
        "method": "deterministic_script_baseline_not_motif",
        "report": str(qmd.with_suffix(".html")) if render else str(qmd),
    }
    (output_dir / f"{scenario}.json").write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")
    return log


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--book", type=Path, default=DEFAULT_BOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--scenario", choices=("base", "duration_update", "result_update"), default="base")
    parser.add_argument("--no-render", action="store_true")
    args = parser.parse_args()
    print(json.dumps(run(args.book, args.output_dir, args.scenario, not args.no_render), ensure_ascii=False, indent=2))
