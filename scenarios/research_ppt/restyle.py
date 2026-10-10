"""候选：受限配色转换，保留布局、正文、原notes与原生媒体/数据。"""
from __future__ import annotations

import hashlib
from pathlib import Path

from .documents import inspect_pptx, safe_zip


# 与 ts/render.ts 对齐；只映射已知模板色，不根据“看起来像绿色”改用户数据。
_THEMES = {
    'academic': {'title': '17345C', 'accent': '117D88', 'body': '243247',
                 'pale': 'F2F6FA', 'border': 'D7DFEA'},
    'lab': {'title': '164A46', 'accent': '207C6B', 'body': '243B39',
            'pale': 'F1F7F4', 'border': 'D5E4DE'},
}
_KNOWN_ACCENTS = {'207C6B', '117D88', '087F8C'}
_KNOWN_HEADERS = {'164A46', '17345C', *_KNOWN_ACCENTS}
_KNOWN_PALE = {'F1F7F4', 'F2F6FA'}
_KNOWN_BORDERS = {'D5E4DE', 'D7DFEA'}


def restyle(source: Path, target: Path, template: str, source_label: str) -> dict:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL_TYPE
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.oxml.ns import qn

    if template not in _THEMES:
        raise ValueError('模板只能是 academic 或 lab')
    theme = _THEMES[template]
    prs = Presentation(str(source))
    before, after, original_notes = [], [], []
    audit = {'title_shapes': [], 'decoration_colors_changed': 0,
             'table_fills_changed': 0, 'dark_cell_text_preserved': 0}

    def rgb(color):
        return str(color.rgb).upper() if color.type == MSO_COLOR_TYPE.RGB else None

    def solid_rgb(fill):
        return rgb(fill.fore_color) if fill.type == MSO_FILL_TYPE.SOLID else None

    def fill_background(fill, parent):
        value = solid_rgb(fill)
        if value is not None:
            return value
        return parent if fill.type in (None, MSO_FILL_TYPE.BACKGROUND) else None

    def set_rgb(color, value):
        color.rgb = RGBColor.from_string(value)

    def luminance(value):
        channels = [int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        channels = [c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
                    for c in channels]
        return sum(c * weight for c, weight in zip(channels, (0.2126, 0.7152, 0.0722)))

    def legible_color(background, preferred):
        if background is None:
            return None  # 未解析的主题/渐变/图片背景不覆盖原字色。
        level, foreground = luminance(background), luminance(preferred)
        preferred_contrast = (max(level, foreground) + 0.05) / (min(level, foreground) + 0.05)
        white_contrast = 1.05 / (level + 0.05)
        return 'FFFFFF' if white_contrast > preferred_contrast else preferred

    def leaf_shapes(shapes):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from leaf_shapes(shape.shapes)
            else:
                yield shape

    def font_size(shape):
        """读直接run/段落/list默认字号，不把未知字号猜成大标题。"""
        sizes = []
        for paragraph in shape.text_frame.paragraphs:
            paragraph_size = paragraph.font.size
            defaults = [paragraph_size.pt] if paragraph_size is not None else []
            end = paragraph._p.find(qn('a:endParaRPr'))
            if end is not None and end.get('sz'):
                defaults.append(int(end.get('sz')) / 100)
            style = shape.text_frame._txBody.find(qn('a:lstStyle'))
            if style is not None:
                level = style.find(qn(f'a:lvl{paragraph.level + 1}pPr'))
                default = None if level is None else level.find(qn('a:defRPr'))
                if default is not None and default.get('sz'):
                    defaults.append(int(default.get('sz')) / 100)
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                if run.font.size is not None:
                    sizes.append(run.font.size.pt)
                elif defaults:
                    sizes.append(max(defaults))
        return max(sizes) if sizes else None

    def find_titles(slide):
        text_shapes = [shape for shape in leaf_shapes(slide.shapes)
                       if shape.has_text_frame and shape.text.strip()]
        placeholders = [shape for shape in text_shapes if shape.is_placeholder and
                        shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]
        if placeholders:
            return {shape.shape_id for shape in placeholders}, 'title_placeholder'
        # PptxGenJS标题是普通文本shape。只有明显最大且在上半页的宽文本框才推断为标题；
        # 同字号正文、底部强调句或未知继承字号不擅自选一个涂色。
        ranked = [(font_size(shape), shape) for shape in text_shapes]
        ranked = sorted([(size, shape) for size, shape in ranked if size is not None],
                        key=lambda item: (-item[0], item[1].top))
        if not ranked:
            return set(), 'no_resolved_font_size'
        largest, candidate = ranked[0]
        next_largest = ranked[1][0] if len(ranked) > 1 else 0
        if (largest < 24 or candidate.top > prs.slide_height * 0.5 or
                candidate.width < prs.slide_width * 0.25 or len(candidate.text) > 200 or
                (next_largest and largest < next_largest * 1.15)):
            return set(), 'ambiguous_nonplaceholder_title'
        return {candidate.shape_id}, 'largest_effective_font'

    def change_decoration(shape):
        # 仅顶层已知模板边框/线段几何；group内的矢量科研图件和图表不改。
        if shape.shape_type not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE):
            return
        if shape.has_text_frame and shape.text.strip():
            return
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
        side_bar = (left <= prs.slide_width * 0.03 and width <= prs.slide_width * 0.03
                    and height >= prs.slide_height * 0.65)
        top_rule = (top <= prs.slide_height * 0.28 and height <= prs.slide_height * 0.035
                    and (width <= prs.slide_width * 0.2 or width >= prs.slide_width * 0.6))
        if not (side_bar or top_rule):
            return
        shape_fill = getattr(shape, 'fill', None)
        for color in (shape_fill.fore_color if shape_fill is not None and shape_fill.type == MSO_FILL_TYPE.SOLID else None,
                      shape.line.color):
            if color is None:
                continue
            old = rgb(color)
            new = theme['accent'] if old in _KNOWN_ACCENTS else theme['border'] if old in _KNOWN_BORDERS else None
            if new is not None and old != new:
                set_rgb(color, new)
                audit['decoration_colors_changed'] += 1

    def visit(shapes, collect, update=False, titles=None, background=None):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit(shape.shapes, collect, update, titles, background)
            if shape.has_text_frame:
                collect.append(shape.text)
                if update and shape.text.strip():
                    is_title = shape.shape_id in titles
                    foreground = legible_color(fill_background(shape.fill, background),
                                               theme['title'] if is_title else theme['body'])
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = 'Noto Sans CJK SC'
                            if foreground is not None:
                                set_rgb(run.font.color, foreground)
                            if is_title:
                                run.font.bold = True
                            # 不补写猜测字号，保留原显式与继承字号、位置、行距。
            if shape.has_table:
                for row_index, row in enumerate(shape.table.rows):
                    for cell in row.cells:
                        collect.append(cell.text)
                        if not update:
                            continue
                        fill = solid_rgb(cell.fill)
                        mapped = (theme['title'] if row_index == 0 and fill in _KNOWN_HEADERS else
                                  theme['pale'] if fill in _KNOWN_PALE else None)
                        if mapped is not None and mapped != fill:
                            set_rgb(cell.fill.fore_color, mapped)
                            fill = mapped
                            audit['table_fills_changed'] += 1
                        foreground = legible_color(fill_background(cell.fill, background), theme['body'])
                        if foreground == 'FFFFFF':
                            audit['dark_cell_text_preserved'] += 1
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.name = 'Noto Sans CJK SC'
                                if foreground is not None:
                                    set_rgb(run.font.color, foreground)

    def geometry(presentation):
        pages = []
        for slide in presentation.slides:
            shapes = []
            for shape in leaf_shapes(slide.shapes):
                item = [shape.shape_id, int(shape.shape_type), shape.left, shape.top,
                        shape.width, shape.height, shape.rotation]
                if shape.has_table:
                    item += [[row.height for row in shape.table.rows],
                             [column.width for column in shape.table.columns]]
                shapes.append(item)
            pages.append(shapes)
        return [presentation.slide_width, presentation.slide_height, pages]

    original_geometry = geometry(prs)
    for number, slide in enumerate(prs.slides, 1):
        visit(slide.shapes, before)
        original_notes.append(slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else '')
        titles, title_method = find_titles(slide)
        audit['title_shapes'].append({'page': number, 'shape_ids': sorted(titles), 'method': title_method})
        background = solid_rgb(slide.background.fill)
        # 不把用户自定义背景整页刷白，也不改图片、渐变、主题填充或图表背景。
        visit(slide.shapes, [], True, titles, background)
        for shape in slide.shapes:
            change_decoration(shape)
        notes = slide.notes_slide.notes_text_frame
        notes.add_paragraph().text = f'来源：{source_label}，原幻灯片 {number}；风格：{template}'
    prs.save(str(target))
    revised = Presentation(str(target))
    for slide in revised.slides:
        visit(slide.shapes, after)
    if before != after:
        raise ValueError('风格转换改变了原始文字，拒绝交付')
    if original_geometry != geometry(revised):
        raise ValueError('风格转换改变了原始布局，拒绝交付')
    for old, slide in zip(original_notes, revised.slides):
        if old and not slide.notes_slide.notes_text_frame.text.startswith(old):
            raise ValueError('风格转换改变了原始讲者备注，拒绝交付')

    def assets(path):
        with safe_zip(path) as archive:
            return {name: hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()
                    if name.startswith(('ppt/media/', 'ppt/charts/', 'ppt/embeddings/')) and not name.endswith('/')}
    if assets(source) != assets(target):
        raise ValueError('风格转换改变了原始图像、图表或嵌入数据，拒绝交付')
    return {**inspect_pptx(target), 'template': template, 'content_preserved': True,
            'layout_preserved': True, 'original_notes_preserved': True,
            'media_charts_preserved': True, 'style_audit': audit,
            'scope': '受限模板配色、明确标题及字体转换；保留原布局/原notes/媒体数据；复杂布局仍需真实渲染验收'}
