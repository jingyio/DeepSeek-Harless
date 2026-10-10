"""真实颜色/保持性反例；不调用模型，不代表科研内容质量。"""
import hashlib
import importlib.util
import os
from pathlib import Path
import tempfile
import unittest
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from scenarios.research_ppt.service import ROOT


def restyler():
    candidate = os.environ.get('SSS_PPT_RESTYLE_CANDIDATE')
    if candidate:
        spec = importlib.util.spec_from_file_location('scenarios.research_ppt.candidate_v4_restyle', candidate)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module.restyle
    from scenarios.research_ppt.restyle import restyle
    return restyle


def text(slide, value, size, y, color='243247'):
    shape = slide.shapes.add_textbox(Inches(.7), Inches(y), Inches(8), Inches(.8))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return shape


class RestyleContrastTests(unittest.TestCase):
    def setUp(self):
        area = ROOT / '.local/research-ppt/restyle-regression'
        area.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=area)
        self.area = Path(self.temp.name)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)

    def tearDown(self):
        self.temp.cleanup()

    def slide(self, background='FFFFFF'):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        return slide

    def output(self):
        source, output = self.area/'source.pptx', self.area/'output.pptx'
        self.prs.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        result = restyler()(source, output, 'academic', '回归来源 SHA256 '+original)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)
        return Presentation(output), result, source, output

    def test_real_dark_table_and_nonplaceholder_title_failure(self):
        slide = self.slide()
        title = text(slide, '原生标题', 32, .4, '164A46')
        text(slide, '实验条件保留', 20, 1.5)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.12), Inches(7.5))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string('207C6B')
        table = slide.shapes.add_table(2, 2, Inches(5), Inches(3), Inches(5), Inches(2)).table
        for index, (row, values) in enumerate(zip(table.rows, [('条件', '指标'), ('独立验证集', '0.8')])):
            for cell, value in zip(row.cells, values):
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string('164A46' if index == 0 else 'FFFFFF')
        for cell in table.rows[0].cells:
            cell.fill.fore_color.rgb = RGBColor.from_string('164A46')
            cell.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        revised, audit, _, _ = self.output()
        shapes = revised.slides[0].shapes
        self.assertEqual(str(shapes[0].text_frame.paragraphs[0].runs[0].font.color.rgb), '17345C')
        self.assertIn(title.shape_id, audit['style_audit']['title_shapes'][0]['shape_ids'])
        self.assertEqual(str(shapes[2].fill.fore_color.rgb), '117D88')
        self.assertTrue(all(str(c.text_frame.paragraphs[0].runs[0].font.color.rgb) == 'FFFFFF' for c in shapes[3].table.rows[0].cells))

    def test_unknown_dark_user_background_is_preserved_and_legible(self):
        slide = self.slide('101010')
        text(slide, '用户深色背景', 32, .4, 'FFFFFF')
        text(slide, '原始正文', 20, 1.5, 'FFFFFF')
        decoration = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.12), Inches(7.5))
        decoration.fill.solid(); decoration.fill.fore_color.rgb = RGBColor.from_string('FF0000')
        revised, _, _, _ = self.output()
        page = revised.slides[0]
        self.assertEqual(str(page.background.fill.fore_color.rgb), '101010')
        self.assertEqual(str(page.shapes[2].fill.fore_color.rgb), 'FF0000')
        for shape in list(page.shapes)[:2]:
            self.assertEqual(str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb), 'FFFFFF')

    def test_equal_size_body_is_not_inferred_as_a_title(self):
        slide = self.slide()
        text(slide, '并列内容甲', 28, .4)
        text(slide, '并列内容乙', 28, 1.5)
        _, result, _, _ = self.output()
        self.assertEqual(result['style_audit']['title_shapes'][0]['shape_ids'], [])

    def test_placeholder_notes_geometry_and_editable_media_are_preserved(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[1])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        slide.shapes.title.text = '原生placeholder标题'
        slide.notes_slide.notes_text_frame.text = '原notes：数值条件与原始来源。'
        picture = self.area/'picture.png'; Image.new('RGB', (200, 100), 'blue').save(picture)
        slide.shapes.add_picture(str(picture), Inches(5), Inches(2), width=Inches(2))
        data = ChartData(); data.categories = ['A', 'B']; data.add_series('原数据', [1.2, 2.4])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(4), Inches(4), Inches(2), data)
        geometry = [(s.shape_id, s.left, s.top, s.width, s.height) for s in slide.shapes]
        revised, result, source, output = self.output()
        self.assertEqual(result['style_audit']['title_shapes'][0]['method'], 'title_placeholder')
        self.assertEqual(geometry, [(s.shape_id, s.left, s.top, s.width, s.height) for s in revised.slides[0].shapes])
        self.assertTrue(revised.slides[0].notes_slide.notes_text_frame.text.startswith('原notes：数值条件与原始来源。'))
        def assets(p):
            with zipfile.ZipFile(p) as z:
                return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()
                        if n.startswith(('ppt/media/', 'ppt/charts/', 'ppt/embeddings/'))}
        self.assertEqual(assets(source), assets(output))


if __name__ == '__main__':
    unittest.main()
