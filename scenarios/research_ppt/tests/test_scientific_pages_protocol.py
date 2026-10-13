"""科研页型的来源绑定和真实服务回归；只在服务器执行并保留外层测试日志。"""
import copy
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from defusedxml import ElementTree as ET

import fitz
from pptx import Presentation

from scenarios.research_ppt.cli import prepare
from scenarios.research_ppt.documents import file_hash
from scenarios.research_ppt.revision import merge_slide_updates, validate_scientific_archetype
from scenarios.research_ppt.server import _safe_failure
from scenarios.research_ppt.service import ROOT, PPTService, readable_source_name
from scenarios.research_ppt.tests.test_revision_service import write_source
from scenarios.research_ppt.validation import NS, _read_package, _speaker_notes_text, plan_expectations, validate_delivery


class ScientificPageProtocolTests(unittest.TestCase):
    def plan(self):
        return {'title': 'Explicit source-bound results', 'slides': [
            {'title': 'Conditions', 'bullets': [],
             'sources': [{'source_id': 'source-a', 'page': 2}, {'source_id': 'source-b', 'page': 1}],
             'archetype': 'result-summary', 'table': [['Method', 'Metric'], ['A', '12']],
             'evidence_notes': [{'label': 'Conditions', 'text': 'Measured on the specified split.', 'source_index': 1}]}]}

    def test_invalid_and_implicit_rebinding_reject_atomically(self):
        original = self.plan()
        before = copy.deepcopy(original)
        sources = list(reversed(original['slides'][0]['sources']))
        with self.assertRaisesRegex(ValueError, '显式更新或移除'):
            merge_slide_updates(original, [{'page': 1, 'changes': {'sources': sources}}])
        notes = [{'label': 'Conditions', 'text': 'Measured on the specified split.', 'source_index': 2}]
        merged = merge_slide_updates(original, [{'page': 1, 'changes': {'sources': sources, 'evidence_notes': notes}}])
        self.assertEqual(merged['slides'][0]['evidence_notes'][0]['source_index'], 2)
        self.assertEqual(original, before)
        for index in (0, 3, True, '1'):
            with self.subTest(index=index), self.assertRaises(ValueError):
                merge_slide_updates(original, [{'page': 1, 'changes': {'evidence_notes': [{**notes[0], 'source_index': index}]}}])
        removed = merge_slide_updates(original, [{'page': 1, 'changes': {'archetype': None, 'evidence_notes': None}}])
        self.assertNotIn('archetype', removed['slides'][0])
        self.assertNotIn('evidence_notes', removed['slides'][0])
        with self.assertRaises(ValueError):
            merge_slide_updates(original, [{'page': 1, 'changes': {'archetype': None}}])

    def test_cross_field_metadata_rejects_unknown_capacity_and_subject_conflicts(self):
        mutations = [
            lambda s: s.update(archetype='evidence-focus'),
            lambda s: s.update(archetype='unknown'),
            lambda s: s.pop('archetype'),
            lambda s: s.update(evidence_notes=None),
            lambda s: s.update(evidence_notes=s['evidence_notes'] * 4),
            lambda s: s['evidence_notes'][0].update(label='x' * 25),
            lambda s: s['evidence_notes'][0].update(text='x' * 111),
            lambda s: s['evidence_notes'][0].update(text='bad\x00text'),
            lambda s: s['evidence_notes'][0].update(extra='unknown'),
        ]
        for mutate in mutations:
            slide = self.plan()['slides'][0]
            mutate(slide)
            with self.subTest(slide=slide), self.assertRaises(ValueError):
                validate_scientific_archetype(slide)
        legacy = {'title': 'Keep', 'bullets': [], 'sources': ['original source']}
        validate_scientific_archetype(legacy)

    def test_source_labels_do_not_guess_authors_or_reveal_directories(self):
        name = 'a' * 80 + '-actual-paper.pdf'
        label = readable_source_name('C:/private/directory/' + name)
        self.assertEqual(len(label), 56)
        self.assertTrue(label.endswith('-actual-paper.pdf'))
        self.assertEqual(readable_source_name(name), label)
        self.assertNotIn('private', label)
        self.assertEqual(readable_source_name('real-paper.pdf'), 'real-paper.pdf')

    def test_notes_body_preserves_explicit_user_text_and_ignores_other_placeholders(self):
        root = ET.fromstring('''<p:notes xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree>
          <p:sp><p:nvSpPr><p:nvPr><p:ph type="hdr"/></p:nvPr></p:nvSpPr><p:txBody>
            <a:p><a:r><a:t>Forged source page 99</a:t></a:r></a:p></p:txBody></p:sp>
          <p:sp><p:nvSpPr><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr><p:txBody>
            <a:p><a:r><a:t>User crop_literal and rebuild discussion</a:t></a:r><a:br/>
              <a:r><a:t>must stay.</a:t></a:r></a:p>
            <a:p><a:r><a:t>source.pdf 第 2 页; SHA256 unchanged</a:t></a:r></a:p></p:txBody></p:sp>
          </p:spTree></p:cSld></p:notes>''')
        actual = _speaker_notes_text(root)
        self.assertEqual(actual, 'User crop_literal and rebuild discussion\nmust stay.\nsource.pdf 第 2 页; SHA256 unchanged')
        self.assertNotIn('99', actual)
        tree = root.find('p:cSld/p:spTree', NS)
        tree.append(copy.deepcopy(tree[-1]))
        with self.assertRaisesRegex(ValueError, '多个讲者正文'):
            _speaker_notes_text(root)

    def test_new_failures_are_public_recoverable_without_sensitive_suffixes(self):
        for message in ('科研页型与主体内容类型不一致', 'evidence_notes 仅允许用于科研页型',
                        'evidence_notes 的 source_index 超出原 sources 范围',
                        '修改 sources 必须显式更新或移除 evidence_notes',
                        '修订字段 evidence_notes 的类型或结构无效'):
            self.assertEqual(_safe_failure(ValueError(message))['error'], 'invalid_plan')
            hidden = _safe_failure(ValueError(message + ' credential=secret'))
            self.assertEqual(hidden['error'], 'internal_processing_failed')
            self.assertNotIn('secret', json.dumps(hidden))


class ScientificPageServiceTests(unittest.TestCase):
    def setUp(self):
        area = ROOT / '.local/research-ppt/scientific-pages-protocol-tests'
        area.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=area)
        self.directory = Path(self.temporary.name)
        self.source_file = self.directory / 'actual-source.pdf'
        with fitz.open() as pdf:
            page = pdf.new_page(width=600, height=800)
            page.insert_text((50, 70), 'The page supplies actual provenance for protocol checks.')
            page.draw_rect(fitz.Rect(100, 150, 400, 300), color=(0, 0, 0))
            page.draw_line((100, 300), (400, 150), color=(0, 0, 0))
            page.insert_text((100, 320), 'Figure 1: Synthetic figure for source-binding diagnostics.')
            second = pdf.new_page(width=600, height=800)
            second.insert_text((50, 70), 'All values in this test are synthetic fixture data.')
            pdf.save(self.source_file)
        self.job = prepare([self.source_file], 'Verify three editable scientific page types.', 3, True)
        self.service = PPTService(self.job)
        self.service.rsi = self.directory / 'rsi'
        self.service.rsi.mkdir()
        self.source = self.service.pin_source('doc_1')['source_id']
        self.service.read_source(self.source)
        candidates = self.service.read_page(self.source, 1)['figure_candidates']
        self.assertEqual(len(candidates), 1)
        image = self.service.extract_figure(self.source, 1, candidate_id=candidates[0]['candidate_id'])
        self.assertEqual(image['evidence_scope'], 'caption_geometry_candidate')
        sources = [{'source_id': self.source, 'page': 1}]
        self.plan = {'title': 'Scientific page protocol fixtures', 'slides': [
            {'title': 'Source figure evidence', 'bullets': [], 'sources': sources, 'archetype': 'evidence-focus',
             'notes': 'User explanation keeps crop_literal and a rebuild comparison; do not infer or delete it.',
             'image': {'source_id': self.source, 'image_id': image['image_id']},
             'evidence_notes': [{'label': 'Scope', 'text': 'Synthetic figure; geometry does not certify interpretation.', 'source_index': 1}]},
            {'title': 'Fixture results', 'bullets': [], 'sources': sources, 'archetype': 'result-summary',
             'table': [['Method', 'Metric'], ['Fixture', '12']],
             'evidence_notes': [{'label': 'Conditions', 'text': 'Synthetic protocol values, not paper measurements.', 'source_index': 1}]},
            {'title': 'Method with explicit explanation', 'bullets': ['Preserve the explicit procedure.'],
             'sources': sources, 'archetype': 'method-map', 'layout': 'process',
             'process': {'steps': [{'label': 'Read', 'detail': 'Read the approved source.'},
                                   {'label': 'Check', 'detail': 'Check the actual generated output.'}]},
             'evidence_notes': [{'label': 'Boundary', 'text': 'The tool does not certify scientific interpretation.', 'source_index': 1}]}]}

    def tearDown(self):
        self.temporary.cleanup()

    def test_forged_visible_metadata_and_invalid_source_index_create_no_output(self):
        paths = set(self.service.output.iterdir()) if self.service.output.exists() else set()
        for mutation in ({'visible_sources': ['Invented Author 2026']},
                         {'evidence_notes': [{'label': 'Scope', 'text': 'Unsupported', 'source_index': 2}]},
                         {'archetype': 'result-summary'}):
            plan = copy.deepcopy(self.plan)
            plan['slides'][0].update(mutation)
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                self.service.render_deck(plan)
            self.assertEqual(set(self.service.output.iterdir()) if self.service.output.exists() else set(), paths)

    def test_real_three_page_render_preserves_visible_provenance_and_full_source_sha_notes(self):
        original = copy.deepcopy(self.plan)
        rendered = self.service.render_deck(self.plan)
        self.assertEqual(self.plan, original)
        saved = json.loads((Path(rendered['path']).parent / 'plan.json').read_text(encoding='utf-8'))
        self.assertEqual(saved['slides'][0]['visible_sources'], ['[1] actual-source.pdf · p.1 · 图注候选 Figure 1'])
        self.assertTrue(all('SHA256' not in label for slide in saved['slides'] for label in slide['visible_sources']))
        self.assertEqual(self.service.read_deck_plan(rendered['deck_id'])['plan'], original)
        sha = file_hash(self.source_file)
        pptx = Presentation(rendered['path'])
        for index, slide in enumerate(pptx.slides):
            note = original['slides'][index]['evidence_notes'][0]
            expected = f"证据说明 [1] {note['label']}：{note['text']}"
            self.assertIn(expected, slide.notes_slide.notes_text_frame.text)
            self.assertIn('SHA256 ' + sha, slide.notes_slide.notes_text_frame.text)
            self.assertIn('; 文档 doc_1', slide.notes_slide.notes_text_frame.text)
        self.assertIn(original['slides'][0]['notes'], pptx.slides[0].notes_slide.notes_text_frame.text)
        figure_id = original['slides'][0]['image']['image_id']
        self.assertNotIn(figure_id, pptx.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('图注候选 Figure 1', pptx.slides[0].notes_slide.notes_text_frame.text)
        metadata = json.loads((Path(rendered['path']).parent / 'source-provenance.json').read_text(encoding='utf-8'))
        self.assertEqual(metadata['pages'][0]['figures'][0]['image_id'], figure_id)
        self.assertEqual(metadata['pages'][0]['figures'][0]['source_sha256'], sha)
        self.assertEqual(metadata['pages'][0]['figures'][0]['figure_label'], 'Figure 1')
        validated = self.service.validate_deck(rendered['deck_id'])
        self.assertTrue(validated['passed'], validated['errors'])
        self.assertEqual(validated['pdf_page_count'], 3)
        # 真实修订也独立检查旧说明是否静默重绑，并核验非目标页原包保持。
        different_source = [{'source_id': self.source, 'page': 2}]
        before_paths = set(self.service.output.iterdir())
        with self.assertRaisesRegex(ValueError, '显式更新或移除'):
            self.service.revise_deck(rendered['deck_id'], [{'page': 2, 'changes': {'sources': different_source}}])
        self.assertEqual(set(self.service.output.iterdir()), before_paths)
        new_note = [{'label': 'Conditions', 'text': 'Fixture values now explicitly cite source page two.', 'source_index': 1}]
        revised = self.service.revise_deck(rendered['deck_id'], [{'page': 2, 'changes': {
            'sources': different_source, 'evidence_notes': new_note}}])
        self.assertTrue(revised['revision_audit']['passed'])
        self.assertTrue(revised['revision_audit']['non_target_slides_preserved'])
        updated_plan = self.service.read_deck_plan(revised['deck_id'])['plan']
        self.assertEqual(updated_plan['slides'][0], original['slides'][0])
        self.assertEqual(updated_plan['slides'][2], original['slides'][2])
        self.assertEqual(updated_plan['slides'][1]['evidence_notes'], new_note)
        fresh = self.service.validate_deck(revised['deck_id'])
        self.assertTrue(fresh['passed'], fresh['errors'])
        self.assertIn('actual-source.pdf 第 2 页; SHA256 ' + sha,
                      Presentation(revised['path']).slides[1].notes_slide.notes_text_frame.text)

    def test_actual_package_notes_use_unique_body_relationship_not_header_or_page_number(self):
        rendered = self.service.render_deck(self.plan)
        source = Path(rendered['path'])
        disguised = self.directory / 'notes-in-header.pptx'
        duplicate = self.directory / 'duplicate-notes-relation.pptx'
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(disguised, 'w') as target:
            changed = False
            for item in original.infolist():
                raw = original.read(item.filename)
                if item.filename == 'ppt/notesSlides/notesSlide1.xml':
                    root = ET.fromstring(raw)
                    bodies = [ph for ph in root.findall('.//p:ph', NS) if ph.get('type') == 'body']
                    self.assertEqual(len(bodies), 1)
                    bodies[0].set('type', 'hdr')
                    raw = ET.tostring(root, encoding='utf-8')
                    changed = True
                target.writestr(item, raw)
            self.assertTrue(changed)
        _, notes = _read_package(disguised, {'errors': [], 'warnings': []})
        self.assertEqual(notes[0], '', '来源文字只在header中不能冒充speaker正文')
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(duplicate, 'w') as target:
            changed = False
            for item in original.infolist():
                raw = original.read(item.filename)
                if item.filename == 'ppt/slides/_rels/slide1.xml.rels':
                    root = ET.fromstring(raw)
                    relations = [rel for rel in root if rel.get('Type', '').endswith('/notesSlide')]
                    self.assertEqual(len(relations), 1)
                    copied = copy.deepcopy(relations[0])
                    copied.set('Id', 'duplicateSpeakerNote')
                    root.append(copied)
                    raw = ET.tostring(root, encoding='utf-8')
                    changed = True
                target.writestr(item, raw)
            self.assertTrue(changed)
        with self.assertRaisesRegex(ValueError, '备注关系不唯一'):
            _read_package(duplicate, {'errors': [], 'warnings': []})
        saved = json.loads((source.parent / 'plan.json').read_text(encoding='utf-8'))
        expected = plan_expectations(saved)
        self.assertIn(self.plan['slides'][0]['notes'], expected['speaker_notes'][0])

    def test_lost_user_speaker_notes_fail_even_when_all_source_labels_are_preserved(self):
        rendered = self.service.render_deck(self.plan)
        source = Path(rendered['path'])
        saved = json.loads((source.parent / 'plan.json').read_text(encoding='utf-8'))
        expected = plan_expectations(saved)
        altered = self.directory / 'user-notes-removed.pptx'
        marker = self.plan['slides'][0]['notes'].encode('utf-8')
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(altered, 'w') as target:
            changed = False
            for item in original.infolist():
                raw = original.read(item.filename)
                if item.filename == 'ppt/notesSlides/notesSlide1.xml':
                    self.assertIn(marker, raw)
                    raw = raw.replace(marker, b'Removed user explanation.')
                    changed = True
                target.writestr(item, raw)
            self.assertTrue(changed)
        report = validate_delivery(altered, 3, expected_texts=expected['texts'],
                                   expected_sources=[slide['sources'] for slide in saved['slides']],
                                   expected_notes=expected['speaker_notes'])
        errors = {row['code'] for row in report['errors']}
        self.assertFalse(report['passed'])
        self.assertIn('speaker_notes_missing_or_changed', errors)
        self.assertNotIn('source_notes_missing', errors)

    def test_real_service_compacts_only_consecutive_source_names_and_restores_interleaved_context(self):
        secondary_file = self.directory / 'secondary-source.pdf'
        write_source(secondary_file, ['A separate source establishes a distinct filename context.',
                                      'Page two of the separate synthetic source.'])
        job = prepare([self.source_file, secondary_file], 'Check compact visible source references.', 3, True)
        service = PPTService(job)
        service.rsi = self.directory / 'compact-source-rsi'
        service.rsi.mkdir()
        primary = service.pin_source('doc_1')['source_id']
        secondary = service.pin_source('doc_2')['source_id']
        service.read_source(primary)
        service.read_source(secondary)
        citations = [{'source_id': primary, 'page': 1}, {'source_id': primary, 'page': 2},
                     {'source_id': secondary, 'page': 1}, {'source_id': secondary, 'page': 2},
                     {'source_id': primary, 'page': 2}, {'source_id': primary, 'page': 1}]
        plan = {'title': 'Compact provenance fixtures', 'slides': [
            {'title': 'Two distinct source contexts', 'bullets': [], 'sources': citations,
             'archetype': 'result-summary', 'table': [['Fixture', 'Value'], ['Synthetic', '12']],
             'evidence_notes': [{'label': 'Association', 'text': 'This explanation cites the second page of the separate source.', 'source_index': 4}]},
            {'title': 'Filename context restarts on each page', 'bullets': [],
             'sources': [{'source_id': primary, 'page': 2}, {'source_id': primary, 'page': 1}],
             'archetype': 'result-summary', 'table': [['Fixture', 'Value'], ['Synthetic', '12']]},
            {'title': 'Legacy page remains unchanged', 'bullets': ['Ordinary source notes retain full evidence.'],
             'sources': citations}]}
        original = copy.deepcopy(plan)
        rendered = service.render_deck(plan)
        saved = json.loads((Path(rendered['path']).parent / 'plan.json').read_text(encoding='utf-8'))
        expected = ['[1] actual-source.pdf · p.1', '[2] p.2',
                    '[3] secondary-source.pdf · p.1', '[4] p.2',
                    '[5] actual-source.pdf · p.2', '[6] p.1']
        self.assertEqual(saved['slides'][0]['visible_sources'], expected)
        self.assertEqual(saved['slides'][1]['visible_sources'], ['[1] actual-source.pdf · p.2', '[2] p.1'])
        self.assertNotIn('visible_sources', saved['slides'][2])
        self.assertEqual(saved['slides'][0]['evidence_notes'][0]['source_index'], 4)
        self.assertEqual(plan, original)
        self.assertEqual(service.read_deck_plan(rendered['deck_id'])['plan'], original)
        pptx = Presentation(rendered['path'])
        for index, slide in enumerate(pptx.slides):
            if index < 2:
                visible_text = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                for citation in saved['slides'][index]['visible_sources']:
                    self.assertIn(citation, visible_text)
            notes = slide.notes_slide.notes_text_frame.text
            for citation in plan['slides'][index]['sources']:
                item = service.source(citation['source_id'])['row']
                self.assertIn(f"{item['name']} 第 {citation['page']} 页; SHA256 {item['version_sha256']}", notes)
        self.assertIn('证据说明 [4] Association：' + plan['slides'][0]['evidence_notes'][0]['text'],
                      pptx.slides[0].notes_slide.notes_text_frame.text)
        validated = service.validate_deck(rendered['deck_id'])
        self.assertTrue(validated['passed'], validated['errors'])


if __name__ == '__main__':
    unittest.main()
