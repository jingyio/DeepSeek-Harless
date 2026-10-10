"""局部修订服务真实集成回归：原稿保留、来源守卫和重新渲染交付。

只由服务器测试入口运行。测试实际生成 PPTX/图表与 LibreOffice PDF/PNG，
不替换渲染、校验或 MCP 操作的返回值；人工注入的损坏仅用于拒绝测试。
"""
import copy
import json
from pathlib import Path
import tempfile
import unittest

from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from scenarios.research_ppt.cli import prepare
from scenarios.research_ppt.documents import file_hash, parse_document
from scenarios.research_ppt.service import ROOT, PPTService


def write_source(path: Path, texts):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
            NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        stream.set_data(f'BT /F1 16 Tf 40 700 Td ({escaped}) Tj ET'.encode('ascii'))
        page[NameObject('/Contents')] = writer._add_object(stream)
    with path.open('wb') as output:
        writer.write(output)


class RevisionServiceTests(unittest.TestCase):
    def setUp(self):
        area = ROOT / '.local/research-ppt/revision-service-tests'
        area.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=area)
        self.directory = Path(self.temporary.name)
        self.primary_file = self.directory / 'primary.pdf'
        self.secondary_file = self.directory / 'secondary.pdf'
        write_source(self.primary_file, ['Original result with measured conditions.',
                                        'The table and chart are independent checks.'])
        write_source(self.secondary_file, ['Updated result is measured on the heldout split.'])
        self.job_path = prepare([self.primary_file, self.secondary_file],
            'Only revise the requested page and retain original delivery evidence.', 3, True)
        self.service = PPTService(self.job_path)
        self.service.rsi = self.directory / 'rsi'
        self.service.rsi.mkdir()
        self.primary = self.service.pin_source('doc_1')['source_id']
        self.service.read_source(self.primary)
        self.plan = {'title': 'Revision provenance checks', 'slides': [
            {'title': 'Original result', 'bullets': ['Measured result with original conditions.'],
             'sources': [{'source_id': self.primary, 'page': 1}], 'notes': 'Original evaluation conditions.'},
            {'title': 'Untouched table', 'bullets': [], 'sources': [{'source_id': self.primary, 'page': 2}],
             'notes': 'Retain the table and its source.', 'table': [['Metric', 'Value'], ['Observed', '12']]},
            {'title': 'Untouched chart', 'bullets': [], 'sources': [{'source_id': self.primary, 'page': 2}],
             'notes': 'Retain editable data and chart conditions.',
             'chart': {'type': 'bar', 'categories': ['A', 'B'],
                       'series': [{'name': 'Measured', 'values': [1.2, 2.4]}]}}]}
        self.original = self.service.render_deck(self.plan)
        self.original_id = self.original['deck_id']
        self.original_path = Path(self.original['path'])
        self.original_bytes = self.original_path.read_bytes()

    def tearDown(self):
        self.temporary.cleanup()

    def updates(self):
        return [{'page': 1, 'changes': {'title': 'Corrected result',
                                      'bullets': ['Corrected result with explicit conditions.']}}]

    def test_explicit_provenance_revision_preserves_original_and_requires_new_real_validation(self):
        old_validation = self.service.validate_deck(self.original_id)
        self.assertTrue(old_validation['passed'], old_validation['errors'])
        original_plan = self.service.read_deck_plan(self.original_id)['plan']
        secondary = self.service.pin_source('doc_2')['source_id']
        self.service.read_source(secondary)
        updates = self.updates()
        updates[0]['changes'].update(sources=[{'source_id': secondary, 'page': 1}],
                                    notes='Updated result is measured on the heldout split.')
        new = self.service.revise_deck(self.original_id, updates)
        new_id, new_path = new['deck_id'], Path(new['path'])
        self.assertNotEqual(new_id, self.original_id)
        self.assertNotEqual(new_path, self.original_path)
        self.assertEqual(new['parent_deck_id'], self.original_id)
        self.assertEqual(new['revision_audit']['checked_pages'], [2, 3])
        self.assertTrue(new['revision_audit']['passed'])
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)
        self.assertEqual(self.service.read_deck_plan(self.original_id)['plan'], original_plan)
        saved = self.service.read_deck_plan(new_id)['plan']
        self.assertEqual(saved['slides'][0]['sources'], [{'source_id': secondary, 'page': 1}])
        self.assertEqual(saved['slides'][1:], original_plan['slides'][1:])
        before, after = Presentation(self.original_path), Presentation(new_path)
        self.assertIn('Updated result is measured on the heldout split.',
                      after.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('secondary.pdf 第 1 页', after.slides[0].notes_slide.notes_text_frame.text)
        for index in (1, 2):
            self.assertEqual(before.slides[index].notes_slide.notes_text_frame.text,
                             after.slides[index].notes_slide.notes_text_frame.text)
        self.assertFalse(any(row['deck_id'] == new_id for row in self.service.validations.values()))
        with self.assertRaisesRegex(ValueError, '未通过真实渲染'):
            self.service.deliver_deck(new_id)
        # An old validation can only deliver its original deck, never authorize the new copy.
        old_delivery = self.service.deliver_deck(old_validation['validation_id'])
        self.assertEqual(old_delivery['deck_id'], self.original_id)
        self.assertEqual(Path(old_delivery['path']), self.original_path)
        self.assertEqual(old_delivery['file_sha256'], file_hash(self.original_path))
        self.assertTrue(self.service.inspect_deck(new_id)['each_slide_has_editable_text'])
        new_validation = self.service.validate_deck(new_id)
        self.assertTrue(new_validation['passed'], new_validation['errors'])
        self.assertNotEqual(new_validation['validation_id'], old_validation['validation_id'])
        self.assertEqual(new_validation['pdf_page_count'], 3)
        previews = [row for row in new_validation['artifacts'] if row['kind'] == 'preview']
        self.assertEqual(len(previews), 3)
        self.assertTrue(all(Path(row['path']).is_file() for row in new_validation['artifacts']))
        self.assertTrue(set(row['path'] for row in old_validation['artifacts']).isdisjoint(
                        row['path'] for row in new_validation['artifacts']))
        delivery = self.service.deliver_deck(new_validation['validation_id'])
        self.assertEqual(delivery['deck_id'], new_id)
        self.assertEqual(delivery['file_sha256'], file_hash(new_path))
        self.assertEqual(delivery['scientific_quality'], '未完成人工事实和视觉审阅')
        # New copy mutations invalidate both plan reads and the fresh validation.
        raw = new_path.read_bytes()
        try:
            new_path.write_bytes(raw + b'\nchanged-after-validation')
            for operation in (lambda: self.service.read_deck_plan(new_id),
                              lambda: self.service.revise_deck(new_id, self.updates()),
                              lambda: self.service.validate_deck(new_id),
                              lambda: self.service.deliver_deck(new_validation['validation_id'])):
                with self.assertRaisesRegex(ValueError, '产物版本已变化'):
                    operation()
        finally:
            new_path.write_bytes(raw)
        self.assertEqual(self.service.deliver_deck(new_validation['validation_id'])['deck_id'], new_id)

    def test_read_plan_is_a_copy_and_omitted_sources_notes_remain_in_actual_pptx(self):
        read = self.service.read_deck_plan(self.original_id)
        self.assertEqual(read['file_sha256'], self.original['file_sha256'])
        read['plan']['slides'][1]['table'][1][1] = '999'
        read['plan']['slides'][0]['sources'][0]['page'] = 999
        self.assertEqual(self.service.read_deck_plan(self.original_id)['plan'], self.plan)
        new = self.service.revise_deck(self.original_id, self.updates())
        plan = self.service.read_deck_plan(new['deck_id'])['plan']
        self.assertEqual(plan['slides'][0]['sources'], self.plan['slides'][0]['sources'])
        self.assertEqual(plan['slides'][0]['notes'], self.plan['slides'][0]['notes'])
        self.assertEqual(plan['slides'][1:], self.plan['slides'][1:])
        before, after = Presentation(self.original_path), Presentation(new['path'])
        self.assertEqual(before.slides[0].notes_slide.notes_text_frame.text,
                         after.slides[0].notes_slide.notes_text_frame.text)
        self.assertIn('Corrected result', parse_document(Path(new['path']))['pages'][0]['text'])
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)

    def test_invalid_page_fields_and_unapproved_unread_or_out_of_range_sources_fail_before_output(self):
        secondary = self.service.pin_source('doc_2')['source_id']  # deliberately unread
        invalid = [[{'page': 4, 'changes': {'title': 'x'}}],
                   [{'page': True, 'changes': {'title': 'x'}}],
                   [{'page': 1, 'changes': {'title': 'x'}}, {'page': 1, 'changes': {'title': 'y'}}],
                   [{'page': 1, 'changes': {'path': '/root/elsewhere'}}],
                   [{'page': 1, 'changes': {'notes': None}}],
                   [{'page': 1, 'changes': {'sources': [{'source_id': 'unapproved', 'page': 1}]}}],
                   [{'page': 1, 'changes': {'sources': [{'source_id': secondary, 'page': 1}]}}],
                   [{'page': 1, 'changes': {'sources': [{'source_id': self.primary, 'page': 999}]}}]]
        directories = set(self.service.output.iterdir())
        for updates in invalid:
            with self.subTest(updates=updates), self.assertRaises(ValueError):
                self.service.revise_deck(self.original_id, updates)
            self.assertEqual(set(self.service.decks), {self.original_id})
            self.assertEqual(set(self.service.output.iterdir()), directories)
            self.assertEqual(self.original_path.read_bytes(), self.original_bytes)

    def test_preview_permission_and_cross_task_deck_access_are_enforced(self):
        self.service.job['allow_output'] = False
        self.assertEqual(self.service.read_deck_plan(self.original_id)['deck_id'], self.original_id)
        with self.assertRaisesRegex(ValueError, '未授权生成文件'):
            self.service.revise_deck(self.original_id, self.updates())
        other = PPTService(prepare([self.primary_file], 'A different task cannot reuse a deck handle.', 3, True))
        with self.assertRaisesRegex(ValueError, '未知 deck_id'):
            other.read_deck_plan(self.original_id)
        with self.assertRaisesRegex(ValueError, '未知 deck_id'):
            other.revise_deck(self.original_id, self.updates())
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)

    def test_source_version_change_blocks_plan_read_and_revision(self):
        write_source(self.primary_file, ['Changed source invalidates the original task.', 'Changed page two.'])
        for operation in (lambda: self.service.read_deck_plan(self.original_id),
                          lambda: self.service.revise_deck(self.original_id, self.updates())):
            with self.assertRaisesRegex(ValueError, '输入版本已变化'):
                operation()
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)

    def test_incorrect_saved_plan_rerender_is_rejected_and_candidate_evidence_is_retained(self):
        # Corrupt the retained request plan, while the checked original PPTX remains intact.
        # The real OOXML guard must catch changes outside the model's requested page.
        self.service.deck_records[self.original_id]['request_plan']['slides'][1]['table'][1][1] = '13'
        with self.assertRaisesRegex(ValueError, '改变了非目标页面'):
            self.service.revise_deck(self.original_id, self.updates())
        self.assertEqual(set(self.service.decks), {self.original_id})
        self.assertEqual(set(self.service.deck_records), {self.original_id})
        self.assertFalse(self.service.validations)
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)
        events = [json.loads(line) for line in (self.job_path.parent / 'feedback.jsonl').read_text(
            encoding='utf-8').splitlines()]
        event = next(row for row in events if row['kind'] == 'revision_rejected')
        candidate = event['candidate_deck_id']
        self.assertTrue((self.service.output / candidate / 'presentation.pptx').is_file())
        self.assertTrue((self.service.output / candidate / 'plan.json').is_file())
        with self.assertRaisesRegex(ValueError, '未知 deck_id'):
            self.service.inspect_deck(candidate)
        with self.assertRaisesRegex(ValueError, '未知 deck_id'):
            self.service.validate_deck(candidate)
        self.assertFalse(any(row['kind'] == 'revision_passed' for row in events))

    def test_restyled_input_without_a_generated_request_plan_cannot_be_partially_revised(self):
        # This test verifies the request-plan boundary, not native-chart restyle
        # compatibility. The chart case previously hit the existing strict asset
        # preservation guard; keep that failure evidence and do not weaken it.
        plain_plan = copy.deepcopy(self.plan)
        plain_plan['slides'][2].pop('chart')
        plain_plan['slides'][2]['bullets'] = ['Editable text for the restyle boundary check.']
        plain = self.service.render_deck(plain_plan)
        restyle = PPTService(prepare([Path(plain['path'])], 'Restyle existing slides.', 3,
                                    True, template='lab', operation='restyle'))
        source = restyle.pin_source('doc_1')['source_id']
        restyle.read_source(source)
        converted = restyle.restyle_deck(source, 'lab')
        for operation in (lambda: restyle.read_deck_plan(converted['deck_id']),
                          lambda: restyle.revise_deck(converted['deck_id'], self.updates())):
            with self.assertRaisesRegex(ValueError, '局部修订仅支持本任务生成'):
                operation()
        self.assertTrue(restyle.inspect_deck(converted['deck_id'])['each_slide_has_editable_text'])
        self.assertEqual(self.original_path.read_bytes(), self.original_bytes)


if __name__ == '__main__':
    unittest.main()
